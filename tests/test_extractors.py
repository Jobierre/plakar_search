import io

import pymupdf
import openpyxl
from docx import Document
from PIL import Image

from plakar_search.extractors import (
    EXTENSIONS,
    SKIP_EXTENSIONS,
    _extract_text,
    _extract_pdf,
    _extract_docx,
    _extract_xlsx,
    _extract_exif,
)


def _make_pdf_bytes(text: str = "Hello PDF") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


def _make_docx_bytes(text: str = "Hello DOCX") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(data: list[list[str]] | None = None) -> bytes:
    if data is None:
        data = [["A1", "B1"], ["A2", "B2"]]
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in data:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_jpeg_bytes(exif_data: dict[int, str] | None = None) -> bytes:
    img = Image.new("RGB", (10, 10), color="red")
    buf = io.BytesIO()
    if exif_data:
        exif = img.getexif()
        for tag_id, value in exif_data.items():
            exif[tag_id] = value
        img.save(buf, format="JPEG", exif=exif.tobytes())
    else:
        img.save(buf, format="JPEG")
    return buf.getvalue()


class TestExtractText:
    def test_plain_text(self):
        result = _extract_text(b"hello world")
        assert result == "hello world"

    def test_utf8_with_accents(self):
        result = _extract_text("cafe".encode("utf-8"))
        assert result == "cafe"

    def test_invalid_utf8_ignored(self):
        result = _extract_text(b"\xff\xfe")
        assert result == ""

    def test_empty_bytes(self):
        result = _extract_text(b"")
        assert result == ""

    def test_registry_text_formats(self):
        for ext in [".txt", ".md", ".csv", ".py", ".js", ".html", ".css",
                     ".json", ".yaml", ".yml", ".xml"]:
            assert ext in EXTENSIONS
            extractor = EXTENSIONS[ext]
            result = extractor(b"test")
            assert result == "test"


class TestExtractPDF:
    def test_single_page(self):
        pdf_bytes = _make_pdf_bytes("Page 1")
        result = _extract_pdf(pdf_bytes)
        assert "Page 1" in result

    def test_multiple_pages(self):
        doc = pymupdf.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {i + 1}")
        pdf_bytes = doc.tobytes()
        result = _extract_pdf(pdf_bytes)
        assert "Page 1" in result
        assert "Page 2" in result
        assert "Page 3" in result

    def test_empty_page(self):
        doc = pymupdf.open()
        doc.new_page()
        pdf_bytes = doc.tobytes()
        result = _extract_pdf(pdf_bytes)
        assert result == ""

    def test_registry_entry(self):
        assert ".pdf" in EXTENSIONS


class TestExtractDOCX:
    def test_single_paragraph(self):
        docx_bytes = _make_docx_bytes("Hello from DOCX")
        result = _extract_docx(docx_bytes)
        assert "Hello from DOCX" in result

    def test_multiple_paragraphs(self):
        doc = Document()
        doc.add_paragraph("Paragraphe 1")
        doc.add_paragraph("Paragraphe 2")
        buf = io.BytesIO()
        doc.save(buf)
        result = _extract_docx(buf.getvalue())
        assert "Paragraphe 1" in result
        assert "Paragraphe 2" in result

    def test_empty_paragraph_skipped(self):
        doc = Document()
        doc.add_paragraph("")  # empty
        doc.add_paragraph("Not empty")
        buf = io.BytesIO()
        doc.save(buf)
        result = _extract_docx(buf.getvalue())
        assert result == "Not empty"

    def test_registry_entry(self):
        assert ".docx" in EXTENSIONS


class TestExtractXLSX:
    def test_basic_data(self):
        xlsx_bytes = _make_xlsx_bytes()
        result = _extract_xlsx(xlsx_bytes)
        assert "A1" in result
        assert "B2" in result

    def test_multiple_sheets(self):
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Feuille1"
        ws1.append(["F1-A", "F1-B"])
        ws2 = wb.create_sheet("Feuille2")
        ws2.append(["F2-A"])

        buf = io.BytesIO()
        wb.save(buf)
        result = _extract_xlsx(buf.getvalue())
        assert "F1-A" in result
        assert "F2-A" in result

    def test_empty_cells_skipped(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, "B1", None])
        buf = io.BytesIO()
        wb.save(buf)
        result = _extract_xlsx(buf.getvalue())
        assert result == "B1"

    def test_registry_entry(self):
        assert ".xlsx" in EXTENSIONS


class TestExtractEXIF:
    def test_jpeg_with_exif(self):
        jpeg_bytes = _make_jpeg_bytes({0x010F: "TestCorp"})  # Make
        result = _extract_exif(jpeg_bytes)
        assert "TestCorp" in result

    def test_jpeg_without_exif(self):
        jpeg_bytes = _make_jpeg_bytes()
        result = _extract_exif(jpeg_bytes)
        assert result == ""

    def test_registry_image_formats(self):
        for ext in [".jpg", ".jpeg", ".png", ".heic"]:
            assert ext in EXTENSIONS


class TestSkipExtensions:
    def test_binary_extensions_present(self):
        assert ".exe" in SKIP_EXTENSIONS
        assert ".zip" in SKIP_EXTENSIONS
        assert ".mp4" in SKIP_EXTENSIONS
        assert ".iso" in SKIP_EXTENSIONS

    def test_binary_not_in_text_registry(self):
        for ext in SKIP_EXTENSIONS:
            assert ext not in EXTENSIONS, f"{ext} ne devrait pas etre dans EXTENSIONS"
