import io

import pymupdf
import openpyxl
from docx import Document
from PIL import Image, ExifTags


def _extract_text(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> str:
    doc = pymupdf.open(stream=data)
    texts = []
    for page in doc:
        t = page.get_text()
        if t:
            texts.append(t)
    doc.close()
    return "\n".join(texts)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    texts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            row_text = " ".join(str(v) for v in row if v is not None)
            if row_text:
                texts.append(row_text)
    wb.close()
    return "\n".join(texts)


def _extract_exif(data: bytes) -> str:
    img = Image.open(io.BytesIO(data))
    exif = img.getexif()
    if not exif:
        return ""

    parts = []
    for tag_id, value in exif.items():
        if tag_id in ExifTags.TAGS:
            tag_name = ExifTags.TAGS[tag_id]
            parts.append(f"{tag_name}: {value}")

    return "\n".join(parts)


EXTENSIONS: dict[str, object] = {
    ".txt": _extract_text,
    ".md": _extract_text,
    ".csv": _extract_text,
    ".py": _extract_text,
    ".js": _extract_text,
    ".html": _extract_text,
    ".css": _extract_text,
    ".json": _extract_text,
    ".yaml": _extract_text,
    ".yml": _extract_text,
    ".xml": _extract_text,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".jpg": _extract_exif,
    ".jpeg": _extract_exif,
    ".png": _extract_exif,
    ".heic": _extract_exif,
}

SKIP_EXTENSIONS: set[str] = {
    ".exe", ".dll", ".so", ".bin",
    ".zip", ".tar", ".gz", ".bz2", ".xz",
    ".7z", ".rar",
    ".mp3", ".mp4", ".avi", ".mov", ".mkv",
    ".ttf", ".otf", ".woff", ".woff2",
    ".iso", ".dmg",
}
